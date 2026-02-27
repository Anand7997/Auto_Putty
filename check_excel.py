import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\VAnand\Downloads\Values.xlsx')
ws = wb.active

print('Columns:', ws.max_column)
print('Rows:', ws.max_row)
print('\nData:')
for i in range(1, min(15, ws.max_row + 1)):
    row_data = [cell.value for cell in ws[i]]
    print(f'Row {i}:', row_data)
